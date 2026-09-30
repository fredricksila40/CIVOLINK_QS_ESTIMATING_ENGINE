"""End-to-end integration (Phase 14): geometry -> rates -> BOQ -> exports.

All figures are ARBITRARY TEST VALUES. Scenario A proves the priced
pipeline conserves every number across layers; B proves fail-closed
behaviour end to end; C proves CESMM lineage on every line.
"""

import ast
import json
from decimal import Decimal
from pathlib import Path

import pytest

from qs_engine.boq.compiler import compile_boq
from qs_engine.boq.measurement import MeasurementRule, MeasurementStandard
from qs_engine.boq.models import (
    Boq,
    BoqItem,
    BoqSection,
    BoqSubsection,
    explicit_rate,
    priced_rate,
    provisional_sum,
    unpriced_item,
)
from qs_engine.contracts.models import Quantity, TakeoffResult
from qs_engine.core.assumptions import VerificationStatus
from qs_engine.core.errors import (
    UnknownMasterDataKeyError,
    UnverifiedMasterDataError,
    ValidationError,
)
from qs_engine.master_data.repository import MasterDataRepository


def repo():
    return MasterDataRepository()


def concrete_overrides():
    return {"dry_factor": "2", "cement_ratio": "1", "sand_ratio": "1",
            "aggregate_ratio": "2", "wc_ratio": "0.5",
            "cement_density_kg_per_m3": "1000",
            "sand_density_kg_per_m3": "1000",
            "aggregate_density_kg_per_m3": "1000",
            "water_density_kg_per_L": "1", "bag_mass_kg": "50"}


def masonry_overrides():
    return {"unit_length_mm": "390", "unit_height_mm": "190",
            "unit_thickness_mm": "200", "bed_joint_mm": "10",
            "perp_joint_mm": "10", "wastage_pct": "5",
            "mortar_cement_ratio": "1", "mortar_sand_ratio": "3",
            "mortar_dry_factor": "1.5",
            "mortar_sand_density_kg_per_m3": "1000",
            "cement_density_kg_per_m3": "1000", "bag_mass_kg": "50"}


def cesmm_rule(ref):
    return MeasurementRule(
        standard=MeasurementStandard("CESMM", "4"), rule_ref=ref,
        source="TEST FIXTURE ONLY",
        verification_status=VerificationStatus.VERIFIED)


def build_rates():
    """Three composite rates (test prices): 242/m3, 500.0/m2, 1050.00/kg."""
    from qs_engine.rates.engine import (
        composite_rate,
        labour_component_direct,
        material_component,
    )
    concrete = composite_rate(
        components=[material_component(
            material_key="portland_cement_42_5", qty_per_unit=Decimal("2"),
            unit_price=Decimal("100"), wastage_pct=Decimal("10"),
            repository=repo())],
        markups=[("profit", Decimal("10"))],
        element="concrete slab", work_unit="m³")
    masonry = composite_rate(
        components=[material_component(
            material_key="concrete_block_200mm",
            qty_per_unit=Decimal("12.5"), unit_price=Decimal("40"),
            repository=repo())],
        markups=[], element="block wall", work_unit="m²")
    steel = composite_rate(
        components=[
            material_component(
                material_key="binding_wire", qty_per_unit=Decimal("1"),
                unit_price=Decimal("700"), repository=repo()),
            labour_component_direct(
                hours_per_unit=Decimal("0.5"), rate_per_hour=Decimal("600"),
                description="fixing"),
        ],
        markups=[("profit", Decimal("5"))],
        element="rebar", work_unit="kg")
    assert concrete.gross_rate == Decimal("242.00")
    assert masonry.gross_rate == Decimal("500.0")
    assert steel.gross_rate == Decimal("1050.00")
    return concrete, masonry, steel


def build_e2e_boq():
    """Slab 5x4x0.15 (3.00 m3), wall 4x2 (8 m2), steel 20x12 m (240 m)."""
    from qs_engine.concrete.engine import (
        ConcreteMixOverride,
        calculate_concrete_requirement,
    )
    from qs_engine.masonry.engine import (
        MasonryUnitOverride,
        calculate_masonry_requirement,
    )
    from qs_engine.reinforcement.engine import (
        BarLineInput,
        RebarBarOverride,
        calculate_bar_line,
        calculate_rebar_schedule,
        resolve_rebar_bar,
    )
    slab = calculate_concrete_requirement(
        length=Decimal("5"), width=Decimal("4"), thickness=Decimal("0.15"),
        class_key="class_c20", repository=repo(),
        overrides=ConcreteMixOverride(**{
            k: Decimal(v) for k, v in concrete_overrides().items()}),
        element="slab S1")
    assert slab.wet_volume.value == Decimal("3.00")
    wall = calculate_masonry_requirement(
        wall_length=Decimal("4"), wall_height=Decimal("2"),
        unit_key="concrete_block_200mm", repository=repo(),
        overrides=MasonryUnitOverride(**{
            k: Decimal(v) for k, v in masonry_overrides().items()}),
        element="wall W1")
    assert wall.takeoff.units_order == 105
    resolved = resolve_rebar_bar(
        "Y12", repo(), RebarBarOverride(mass_kg_per_m=Decimal("2")))
    line = calculate_bar_line(BarLineInput(
        bar_key="Y12", count=20, length_per_bar_m=Decimal("12")), resolved)
    schedule = calculate_rebar_schedule([line], wastage_pct=Decimal("0"))
    assert schedule.total_length_m == Decimal("240")
    assert schedule.total_mass_kg == Decimal("480")

    concrete_rate, masonry_rate, steel_rate = build_rates()
    return compile_boq(Boq(
        project_ref="E2E-PROJ", boq_ref="E2E-BOQ",
        sections=(
            BoqSection(code="C", title="Concrete", subsections=(
                BoqSubsection(code="C.1", title="In-situ", items=(
                    BoqItem(
                        description="Slab concrete", trade="Concrete",
                        unit="m³", quantity=slab.wet_volume.value,
                        pricing=priced_rate(concrete_rate, "R-C1"),
                        takeoff=TakeoffResult(
                            element="slab S1",
                            quantity=Quantity(slab.wet_volume.value, "m³"),
                            method="LxWxT", assumptions=("e2e",)),
                        takeoff_ref="slab S1 take-off",
                        rate_result=concrete_rate,
                        measurement=cesmm_rule("E.1")),
                )),
            )),
            BoqSection(code="M", title="Masonry", subsections=(
                BoqSubsection(code="M.1", title="Blockwork", items=(
                    BoqItem(
                        description="Block wall", trade="Masonry",
                        unit="m²", quantity=wall.wall.net_m2,
                        pricing=priced_rate(masonry_rate, "R-M1"),
                        takeoff=TakeoffResult(
                            element="wall W1",
                            quantity=Quantity(wall.wall.net_m2, "m²"),
                            method="LxH-joints", assumptions=("e2e",)),
                        takeoff_ref="wall W1 take-off",
                        rate_result=masonry_rate,
                        measurement=cesmm_rule("F.2")),
                )),
            )),
            BoqSection(code="R", title="Reinforcement", subsections=(
                BoqSubsection(code="R.1", title="Bars", items=(
                    BoqItem(
                        description="Y12 bars", trade="Reinforcement",
                        unit="kg", quantity=schedule.total_mass_kg,
                        pricing=priced_rate(steel_rate, "R-R1"),
                        takeoff=TakeoffResult(
                            element="rebar S1",
                            quantity=Quantity(
                                schedule.total_mass_kg, "kg"),
                            method="schedule", assumptions=("e2e",)),
                        takeoff_ref="rebar S1 take-off",
                        rate_result=steel_rate,
                        measurement=cesmm_rule("G.3")),
                )),
            )),
        )))


# --- Scenario A: priced workflow --------------------------------------------

def test_a_takeoff_quantities_exact():
    compiled = build_e2e_boq()
    by_no = {i.item_no: i for s in compiled.sections for i in s.items}
    assert by_no["1.1.1"].quantity == Decimal("3.00")
    assert by_no["2.1.1"].quantity == Decimal("8")
    assert by_no["3.1.1"].quantity == Decimal("480")


def test_a_amounts_are_quantity_times_gross():
    compiled = build_e2e_boq()
    by_no = {i.item_no: i for s in compiled.sections for i in s.items}
    assert by_no["1.1.1"].amount == Decimal("726.00")
    assert by_no["2.1.1"].amount == Decimal("4000.0")
    assert by_no["3.1.1"].amount == Decimal("504000.00")
    assert compiled.total == Decimal("508726.00")
    assert compiled.has_unresolved is False


def test_a_excel_reconciliation_matches_compiled(tmp_path):
    from openpyxl import load_workbook
    from qs_engine.adapters.excel.boq_workbook import export_boq_to_excel
    compiled = build_e2e_boq()
    out = export_boq_to_excel(compiled, str(tmp_path / "e2e.xlsx"))
    wb = load_workbook(out, data_only=False)
    assert wb.sheetnames == ["Cover", "Summary", "C", "M", "R"]
    ws = wb["C"]
    assert ws.cell(row=4, column=6).value == "=C4*E4"
    assert ws.cell(row=4, column=7).value == float(Decimal("726.00"))
    assert ws.cell(row=4, column=8).value == "=F4-G4"
    assert wb["Summary"].cell(row=5, column=4).value == float(
        Decimal("504000.0"))
    assert wb["Summary"].cell(row=6, column=4).value == float(
        Decimal("508726.00"))


def test_a_document_payload_matches_compiled():
    import json
    from qs_engine.adapters.document.payload import synthesize_boq_payload
    compiled = build_e2e_boq()
    payload = synthesize_boq_payload(
        compiled, document_title="Tender BOQ Schedule",
        purpose="E2E test", compilation_id="E2E-001")
    assert payload.summary.grand_total == "508726.00"
    assert payload.summary.has_unresolved is False
    data = json.loads(payload.to_json())
    assert Decimal(data["sections"][0]["items"][0]["amount"]) == Decimal(
        "726.00")
    assert data["summary"]["unresolved_items"] == []


def test_a_ui_controllers_agree_with_engine():
    from qs_engine.ui.controllers import (
        boq_view_controller,
        concrete_controller,
        document_payload_controller,
        excel_export_controller,
    )
    view = concrete_controller(
        {"length": "5", "width": "4", "thickness": "0.15",
         "class_key": "class_c20", "element": "slab S1"},
        repo(), concrete_overrides())
    assert view.wet_m3 == "3.00"
    assert view.bags_order == "30"
    boq = Boq(project_ref="E2E", boq_ref="B", sections=(
        BoqSection.single("C", "Concrete", (
                BoqItem(description="Slab", trade="Concrete", unit="m³",
                        quantity=Decimal("3.00"),
                        pricing=explicit_rate(Decimal("242.00"), "R-C1")),
        )),
    ))
    rendered = boq_view_controller(boq)
    assert rendered.grand_total == "726.0000"
    payload = document_payload_controller(
        compile_boq(boq), "T", "P")
    assert payload.summary.grand_total == "726.0000"


# --- Scenario B: fail-closed --------------------------------------------------

def test_b_fail_closed_unverified_class():
    from qs_engine.concrete.engine import resolve_mix
    from qs_engine.core.errors import UnverifiedMasterDataError
    with pytest.raises(UnverifiedMasterDataError):
        resolve_mix("class_c20", repo())


def test_b_ui_controller_reports_errorview():
    from qs_engine.ui.controllers import concrete_controller
    from qs_engine.ui.viewmodels import ErrorView
    view = concrete_controller(
        {"length": "5", "width": "4", "thickness": "0.15",
         "class_key": "class_c20"}, repo())
    assert isinstance(view, ErrorView)
    assert "NEEDS_VERIFICATION" in view.message


def test_b_unpriced_boq_totals_and_views():
    from qs_engine.ui.controllers import boq_view_controller
    boq = Boq(project_ref="E2E", boq_ref="B", sections=(
        BoqSection.single("C", "Concrete", (
            BoqItem(description="No rate yet", trade="Concrete", unit="m³",
                    quantity=Decimal("3"), pricing=unpriced_item()),
        )),
    ))
    compiled = compile_boq(boq)
    assert compiled.total is None
    view = boq_view_controller(boq)
    assert view.grand_total == "—"
    assert view.has_unresolved is True
    assert view.unresolved_items == ("1.1.1",)


def test_b_unknown_material_fails_closed_engine_and_ui():
    from qs_engine.rates.engine import material_component
    with pytest.raises(UnknownMasterDataKeyError):
        material_component(material_key="unobtainium",
                           qty_per_unit=Decimal("1"),
                           unit_price=Decimal("10"), repository=repo())
    from qs_engine.ui.controllers import rate_builder_controller
    from qs_engine.ui.viewmodels import ErrorView
    view = rate_builder_controller(
        [{"kind": "material", "material_key": "unobtainium",
          "qty_per_unit": "1", "unit_price": "10"}],
        [], "sequential", "r", "m³", repo())
    assert isinstance(view, ErrorView)


# --- Scenario C: CESMM lineage --------------------------------------------------

def test_c_every_line_fully_traced():
    compiled = build_e2e_boq()
    for section in compiled.sections:
        for item in section.items:
            assert item.takeoff_ref, item.item_no
            assert item.rate_ref, item.item_no
            assert item.measurement_ref, item.item_no
            assert item.takeoff is not None
            assert item.rate_result is not None
            assert item.takeoff.element in item.takeoff_ref
            assert item.rate_result.gross_rate == item.rate


def test_c_dto_and_excel_carry_citations(tmp_path):
    import json
    from openpyxl import load_workbook
    from qs_engine.adapters.document.payload import synthesize_boq_payload
    from qs_engine.adapters.excel.boq_workbook import export_boq_to_excel
    compiled = build_e2e_boq()
    payload = synthesize_boq_payload(
        compiled, document_title="T", purpose="P")
    assert payload.measurement_standards == (
        "CESMM 4: E.1", "CESMM 4: F.2", "CESMM 4: G.3")
    for section in payload.sections:
        for item in section.items:
            assert len(item.provenance) == 3  # take-off, rate, rule
    out = export_boq_to_excel(compiled, str(tmp_path / "e2e.xlsx"))
    wb = load_workbook(out, data_only=False)
    cover_rules = wb["Cover"].cell(row=12, column=2).value
    assert "CESMM 4: E.1" in cover_rules


# --- architecture invariants -----------------------------------------------------

def test_no_float_in_product_code():
    """Binary floating point must not appear anywhere in qs_engine/."""
    import ast as _ast
    root = Path(__file__).parent.parent / "qs_engine"
    offenders = []
    for path in sorted(root.rglob("*.py")):
        tree = _ast.parse(path.read_text(encoding="utf-8"),
                          filename=str(path))
        for node in _ast.walk(tree):
            if isinstance(node, _ast.Constant) and isinstance(
                    node.value, float):
                offenders.append(f"{path.name}:{node.lineno}")
            if (isinstance(node, _ast.Call)
                    and isinstance(node.func, _ast.Name)
                    and node.func.id == "float"):
                offenders.append(f"{path.name}:{node.lineno}")
    assert offenders == [], f"float usage in product code: {offenders}"
