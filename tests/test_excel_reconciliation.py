"""Phase 17 Excel tender export + reconciliation tests.

Covers the tender workbook layout (header block, live formulas,
unpriced discipline, group summaries), the variance engine (exact
Decimal deltas, missing-rate handling, skip counts), the controller
triggers, and the dual-panel excel page wiring. The Phase 11 legacy
exporter is untouched (see test_excel_adapter.py).
"""

import json
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from qs_engine.adapters.excel.reconciliation import (
    reconcile,
    reconcile_tender_file,
)
from qs_engine.adapters.excel.tender_workbook import (
    TENDER_HEADERS,
    TENDER_STYLES,
    export_tender_to_excel,
)
from qs_engine.boq.compiler import compile_boq
from qs_engine.boq.models import (
    Boq,
    BoqItem,
    BoqSection,
    BoqSubsection,
    explicit_rate,
    provisional_sum,
    unpriced_item,
)
from qs_engine.core.errors import ValidationError
from qs_engine.ui import controllers, views


def sample_compiled():
    boq = Boq(
        project_ref="P17", boq_ref="B17",
        sections=(
            BoqSection(code="E", title="Earthworks", subsections=(
                BoqSubsection(code="E.1", title="Excavation", items=(
                    BoqItem(description="Bulk excavation", trade="Earthworks",
                            unit="m\u00b3", quantity=Decimal("10"),
                            pricing=explicit_rate(Decimal("100"), "test")),
                    BoqItem(description="Dewatering", trade="Earthworks",
                            unit="sum", quantity=Decimal("1"),
                            pricing=provisional_sum(Decimal("500"), "pc")),
                )),
            )),
            BoqSection(code="C", title="Concrete", subsections=(
                BoqSubsection(code="C.1", title="In-situ", items=(
                    BoqItem(description="Blinding", trade="Concrete",
                            unit="m\u00b3", quantity=Decimal("5"),
                            pricing=explicit_rate(Decimal("200"), "test")),
                    BoqItem(description="Awaiting rate", trade="Concrete",
                            unit="m\u00b3", quantity=Decimal("3"),
                            pricing=unpriced_item()),
                )),
            )),
        ))
    return compile_boq(boq)


def priced_compiled():
    boq = Boq(
        project_ref="P17", boq_ref="B17P",
        sections=(
            BoqSection(code="C", title="Concrete", subsections=(
                BoqSubsection(code="C.1", title="In-situ", items=(
                    BoqItem(description="Blinding", trade="Concrete",
                            unit="m\u00b3", quantity=Decimal("5"),
                            pricing=explicit_rate(Decimal("200"), "test")),
                )),
            )),
        ))
    return compile_boq(boq)


def _sheet(path):
    return load_workbook(str(path))["Tender Schedule"]


def _item_row(ws, item_no):
    for row in range(1, ws.max_row + 1):
        if str(ws.cell(row=row, column=1).value or "") == item_no:
            return row
    raise AssertionError(f"item {item_no} not found in tender sheet")


def _retender(path, target, edits):
    """Copy an exported tender file with Quantity/Rate cell edits.

    edits: {item_no: (quantity_or_None, rate_or_None)}; None clears.
    """
    from openpyxl import load_workbook as _load
    wb = _load(str(path))
    ws = wb["Tender Schedule"]
    for item_no, (quantity, rate) in edits.items():
        row = _item_row(ws, item_no)
        ws.cell(row=row, column=4).value = quantity
        ws.cell(row=row, column=5).value = rate
    wb.save(str(target))
    return str(target)


def _sum_cells(ws, formula):
    """Every cell a =SUM(...) formula references (ranges expanded)."""
    inner = formula[formula.index("(") + 1:formula.rindex(")")]
    cells = []
    for ref in inner.split(","):
        got = ws[ref.strip()]
        if isinstance(got, tuple):
            for row in got:
                cells.extend(row if isinstance(row, tuple) else (row,))
        else:
            cells.append(got)
    return cells


# --- tender workbook layout ---------------------------------------------------------

def test_tender_headers_and_block(tmp_path):
    out = export_tender_to_excel(
        sample_compiled(), tmp_path / "tender.xlsx",
        company_name="Test Corp Ltd", project_title="Demo Towers",
        client="ACME", reference="B17",
        measurement_standard="CESMM4")
    ws = _sheet(out)
    assert ws["A1"].value == "TENDER SCHEDULE \u2014 BILL OF QUANTITIES"
    header = None
    for row in range(1, 30):
        if str(ws.cell(row=row, column=1).value or "") == "Item No":
            header = row
            break
    assert header is not None
    assert [ws.cell(row=header, column=c).value for c in range(1, 8)] == (
        list(TENDER_HEADERS))
    block = {}
    for row in range(1, header):
        label = str(ws.cell(row=row, column=1).value or "")
        if label in ("Company Name", "Project Title", "Client",
                     "Reference", "Measurement Standard", "Date"):
            block[label] = ws.cell(row=row, column=2).value
    assert block == {"Company Name": "Test Corp Ltd",
                     "Project Title": "Demo Towers", "Client": "ACME",
                     "Reference": "B17", "Measurement Standard": "CESMM4",
                     "Date": block["Date"]}
    assert str(block["Date"]).strip()  # export stamps a real date


def test_template_styles_and_rejection(tmp_path):
    std = export_tender_to_excel(
        sample_compiled(), tmp_path / "std.xlsx", template_style="standard")
    plain = export_tender_to_excel(
        sample_compiled(), tmp_path / "plain.xlsx",
        template_style="plain")
    header_row = None
    ws = _sheet(std)
    for row in range(1, 30):
        if str(ws.cell(row=row, column=1).value or "") == "Item No":
            header_row = row
            break
    assert str(ws.cell(row=header_row, column=1).fill.start_color.rgb
               ).upper().endswith("004F9D")
    ws_plain = _sheet(plain)
    assert str(ws_plain.cell(row=header_row, column=1).fill.start_color.rgb
               ).upper().endswith("BFBFBF")
    assert set(TENDER_STYLES) == {"standard", "plain"}
    with pytest.raises(ValidationError):
        export_tender_to_excel(sample_compiled(), tmp_path / "x.xlsx",
                               template_style="glossy")
    with pytest.raises(ValidationError):
        export_tender_to_excel("not-a-boq", tmp_path / "x.xlsx")


def test_priced_row_carries_live_formula(tmp_path):
    ws = _sheet(export_tender_to_excel(
        sample_compiled(), tmp_path / "tender.xlsx"))
    row = _item_row(ws, "1.1.1")
    # Column data must align with TENDER_HEADERS so =D*E is valid:
    # C=Unit, D=Quantity, E=Rate (guards the C/D swap regression).
    assert ws.cell(row=row, column=3).value == "m\u00b3"
    assert ws.cell(row=row, column=4).value == 10
    assert ws.cell(row=row, column=5).value == 100
    assert ws.cell(row=row, column=6).value == f"=D{row}*E{row}"
    assert ws.cell(row=row, column=7).value in ("", None)


def test_unpriced_row_blank_never_zero(tmp_path):
    ws = _sheet(export_tender_to_excel(
        sample_compiled(), tmp_path / "tender.xlsx"))
    row = _item_row(ws, "2.1.2")
    assert ws.cell(row=row, column=5).value is None
    assert ws.cell(row=row, column=6).value is None
    assert ws.cell(row=row, column=7).value == "Unpriced Item"


def test_provisional_row_lump_value(tmp_path):
    ws = _sheet(export_tender_to_excel(
        sample_compiled(), tmp_path / "tender.xlsx"))
    row = _item_row(ws, "1.1.2")
    assert ws.cell(row=row, column=5).value is None
    assert ws.cell(row=row, column=6).value == 500
    assert ws.cell(row=row, column=7).value == "Provisional Sum"


def test_section_total_formula_or_unresolved(tmp_path):
    ws = _sheet(export_tender_to_excel(
        sample_compiled(), tmp_path / "tender.xlsx"))
    totals = {}
    for row in range(1, ws.max_row + 1):
        if str(ws.cell(row=row, column=1).value or "") == "SECTION TOTAL":
            section = str(
                ws.cell(row=row - 3, column=1).value or "")[:1]
            # Money lives in the Amount column (F); Status (G) stays
            # empty on total rows — matches section_total_cells refs.
            totals[section] = ws.cell(row=row, column=6).value
            assert ws.cell(row=row, column=7).value in (None, "")
    assert str(totals["E"]).startswith("=SUM(")
    # Every referenced cell must hold a value — an =SUM over blanks
    # silently renders 0.00 (the F/G column-mismatch bug class).
    for cell in _sum_cells(ws, totals["E"]):
        assert cell.value not in (None, ""), (
            f"{totals['E']} references empty {cell.coordinate}")
    assert totals["C"] == "UNRESOLVED"


def test_collection_summary_net_total(tmp_path):
    ws = _sheet(export_tender_to_excel(
        sample_compiled(), tmp_path / "tender.xlsx"))
    labels = {}
    for row in range(1, ws.max_row + 1):
        label = str(ws.cell(row=row, column=1).value or "")
        if label in ("Direct Materials", "Direct Labour",
                     "Plant & Haulage", "Other Sections", "Overheads",
                     "NET TENDER TOTAL"):
            labels[label] = ws.cell(row=row, column=6).value
            assert ws.cell(row=row, column=7).value in (None, "")
    assert labels["Direct Materials"] == "UNRESOLVED"  # C section open
    assert labels["Overheads"] == "\u2014"
    assert labels["NET TENDER TOTAL"] == "UNRESOLVED"
    assert labels["Direct Labour"] == "\u2014"  # no L sections at all
    ws_priced = _sheet(export_tender_to_excel(
        priced_compiled(), tmp_path / "priced.xlsx"))
    priced_vals = {}
    for row in range(1, ws_priced.max_row + 1):
        label = str(ws_priced.cell(row=row, column=1).value or "")
        if label in ("Direct Materials", "NET TENDER TOTAL"):
            priced_vals[label] = ws_priced.cell(row=row, column=6).value
    # Group and net formulas must reference section-total formula
    # cells (never empty F cells, which would sum to 0.00).
    for key in ("Direct Materials", "NET TENDER TOTAL"):
        formula = priced_vals[key]
        assert str(formula).startswith("=SUM("), (key, formula)
        for cell in _sum_cells(ws_priced, formula):
            assert str(cell.value or "").startswith("=SUM("), (
                f"{key} {formula} ref {cell.coordinate} = {cell.value!r}")


# --- reconciliation engine -------------------------------------------------------------

def test_round_trip_no_changes(tmp_path):
    base = tmp_path / "tender.xlsx"
    export_tender_to_excel(sample_compiled(), base)
    summary = reconcile_tender_file(sample_compiled(), base)
    assert summary.base_estimate == "2500"
    assert summary.tender_total == "2500"
    assert summary.tender_has_unpriced is True  # 2.1.2 still open
    assert summary.variance_amount == "0"
    assert summary.variance_pct == "0.00"
    assert summary.modified_count == 0
    assert summary.unpriced_count == 1
    assert all(row.variance in ("0", None) for row in summary.rows)
    priced = [row for row in summary.rows if row.item_no == "1.1.1"][0]
    assert (priced.q_boq, priced.q_tender, priced.dq) == ("10", "10", "0")
    assert priced.variance == "0" and priced.modified is False


def test_rate_change_variance_exact(tmp_path):
    base = tmp_path / "tender.xlsx"
    export_tender_to_excel(sample_compiled(), base)
    returned = _retender(base, tmp_path / "returned.xlsx",
                         {"1.1.1": (10, 110)})
    summary = reconcile_tender_file(sample_compiled(), returned)
    row = [r for r in summary.rows if r.item_no == "1.1.1"][0]
    assert (row.r_boq, row.r_tender, row.dr) == ("100", "110", "10")
    assert row.variance == "100" and row.modified is True
    assert "rate-changed" in row.flags
    assert summary.tender_total == "2600"
    assert summary.variance_amount == "100"
    assert summary.variance_pct == "4.00"
    assert summary.modified_count == 1
    assert summary.rate_changed_count == 1


def test_quantity_change_variance_exact(tmp_path):
    base = tmp_path / "tender.xlsx"
    export_tender_to_excel(sample_compiled(), base)
    returned = _retender(base, tmp_path / "returned.xlsx",
                         {"2.1.1": (6, 200)})
    summary = reconcile_tender_file(sample_compiled(), returned)
    row = [r for r in summary.rows if r.item_no == "2.1.1"][0]
    assert (row.q_boq, row.q_tender, row.dq) == ("5", "6", "1")
    assert row.variance == "200"
    assert "quantity-changed" in row.flags
    assert summary.quantity_changed_count == 1


def test_tender_unpriced_rate_flagged_never_zero(tmp_path):
    base = tmp_path / "tender.xlsx"
    export_tender_to_excel(sample_compiled(), base)
    returned = _retender(base, tmp_path / "returned.xlsx",
                         {"1.1.1": (10, None)})
    summary = reconcile_tender_file(sample_compiled(), returned)
    row = [r for r in summary.rows if r.item_no == "1.1.1"][0]
    assert row.r_tender is None and row.a_tender is None
    assert row.variance is None
    assert "became-unpriced" in row.flags and row.modified is True
    assert summary.tender_has_unpriced is True
    assert summary.tender_total == "1500"  # 500 + 1000 priced only
    assert summary.variance_amount == "-1000"


def test_provisional_lump_change_tracked(tmp_path):
    from openpyxl import load_workbook as _load
    base = tmp_path / "tender.xlsx"
    export_tender_to_excel(sample_compiled(), base)
    wb = _load(str(base))
    ws = wb["Tender Schedule"]
    for row in range(1, ws.max_row + 1):
        if str(ws.cell(row=row, column=1).value or "") == "1.1.2":
            ws.cell(row=row, column=6, value=600)
    revised = tmp_path / "revised.xlsx"
    wb.save(str(revised))
    summary = reconcile_tender_file(sample_compiled(), revised)
    row = [r for r in summary.rows if r.item_no == "1.1.2"][0]
    assert (row.a_boq, row.a_tender, row.variance) == ("500", "600", "100")
    assert "lump-changed" in row.flags and row.modified is True
    assert summary.tender_total == "2600"


def test_base_unpriced_tender_priced(tmp_path):
    base = tmp_path / "tender.xlsx"
    export_tender_to_excel(sample_compiled(), base)
    returned = _retender(base, tmp_path / "returned.xlsx",
                         {"2.1.2": (3, 50)})
    summary = reconcile_tender_file(sample_compiled(), returned)
    row = [r for r in summary.rows if r.item_no == "2.1.2"][0]
    assert row.a_tender == "150" and row.variance is None
    assert "became-priced" in row.flags and row.modified is True


def test_unmatched_rows_counted_and_empty_rejected(tmp_path):
    from openpyxl import load_workbook as _load
    base = tmp_path / "tender.xlsx"
    export_tender_to_excel(sample_compiled(), base)
    wb = _load(str(base))
    ws = wb["Tender Schedule"]
    ws.cell(row=ws.max_row + 2, column=1, value="9.9.9")
    ws.cell(row=ws.max_row, column=4, value=7)
    odd = tmp_path / "odd.xlsx"
    wb.save(str(odd))
    summary = reconcile_tender_file(sample_compiled(), odd)
    assert summary.unmatched_count == 1
    assert summary.modified_count == 0
    renamed = tmp_path / "renamed.xlsx"
    wb = _load(str(base))
    ws = wb["Tender Schedule"]
    for row in range(1, ws.max_row + 1):
        if str(ws.cell(row=row, column=1).value or "") == "1.1.1":
            ws.cell(row=row, column=1, value="9.9.9")
    wb.save(str(renamed))
    with pytest.raises(ValidationError):
        reconcile_tender_file(priced_compiled(), renamed)


def test_zero_base_pct_stays_none(tmp_path):
    boq = Boq(
        project_ref="P0", boq_ref="B0",
        sections=(
            BoqSection(code="C", title="Concrete", subsections=(
                BoqSubsection(code="C.1", title="In-situ", items=(
                    BoqItem(description="Awaiting rate", trade="Concrete",
                            unit="m\u00b3", quantity=Decimal("3"),
                            pricing=unpriced_item()),
                )),
            )),
        ))
    from qs_engine.boq.compiler import compile_boq as _compile
    compiled = _compile(boq)
    path = export_tender_to_excel(compiled, tmp_path / "open.xlsx")
    summary = reconcile_tender_file(compiled, path)
    assert summary.base_estimate == "0"
    assert summary.tender_total == "0"
    assert summary.variance_pct is None


def _resolved_rule():
    from qs_engine.boq.measurement import (
        MeasurementRule,
        MeasurementStandard,
    )
    from qs_engine.core.assumptions import VerificationStatus
    return MeasurementRule(
        standard=MeasurementStandard(name="CESMM", edition="4"),
        rule_ref="5.2",
        verification_status=VerificationStatus.VERIFIED)


def unresolved_compiled():
    from qs_engine.boq.compiler import compile_boq as _compile
    boq = Boq(
        project_ref="PU", boq_ref="BU",
        sections=(
            BoqSection(code="C", title="Concrete", subsections=(
                BoqSubsection(code="C.1", title="In-situ", items=(
                    BoqItem(description="Slab", trade="Concrete",
                            unit="m\u00b3", quantity=Decimal("4"),
                            pricing=explicit_rate(Decimal("100"), "test"),
                            measurement=_resolved_rule()),
                    BoqItem(description="Beams", trade="Concrete",
                            unit="m\u00b3", quantity=Decimal("6"),
                            pricing=explicit_rate(Decimal("50"), "test")),
                )),
            )),
        ))
    return _compile(boq, require_measurement_rules=True)


def test_unresolved_export_shows_rate_withholds_amount(tmp_path):
    ws = _sheet(export_tender_to_excel(
        unresolved_compiled(), tmp_path / "t.xlsx"))
    row = _item_row(ws, "1.1.2")
    assert ws.cell(row=row, column=5).value == 50
    assert ws.cell(row=row, column=6).value is None
    assert ws.cell(row=row, column=7).value == "Unresolved"


def test_unresolved_round_trip_no_spurious_variance(tmp_path):
    compiled = unresolved_compiled()
    path = export_tender_to_excel(compiled, tmp_path / "t.xlsx")
    summary = reconcile_tender_file(compiled, path)
    held = [row for row in summary.rows
            if row.item_no == "1.1.2"][0]
    assert held.a_tender is None and held.variance is None
    assert held.modified is False
    assert "unresolved-open" in held.flags
    assert summary.base_estimate == "400"
    assert summary.tender_total == "400"
    assert summary.variance_amount == "0"
    assert summary.variance_pct == "0.00"
    assert summary.modified_count == 0


def test_unresolved_rate_change_never_inflates_total(tmp_path):
    compiled = unresolved_compiled()
    base = tmp_path / "t.xlsx"
    export_tender_to_excel(compiled, base)
    returned = _retender(base, tmp_path / "returned.xlsx",
                         {"1.1.2": (6, 55)})
    summary = reconcile_tender_file(compiled, returned)
    held = [row for row in summary.rows if row.item_no == "1.1.2"][0]
    assert held.dr == "5" and held.modified is True
    assert "rate-changed" in held.flags
    assert "unresolved-open" in held.flags
    assert held.variance is None and held.a_tender is None
    assert summary.tender_total == "400"
    assert summary.variance_amount == "0"


def test_summary_round_trips_json():
    import tempfile
    from pathlib import Path as _Path
    with tempfile.TemporaryDirectory() as tmp:
        path = _Path(tmp) / "tender.xlsx"
        export_tender_to_excel(sample_compiled(), path)
        summary = reconcile_tender_file(sample_compiled(), path)
    data = json.loads(json.dumps(summary.to_dict()))
    assert data["base_estimate"] == "2500"
    assert Decimal(data["tender_total"]) == Decimal("2500")
    assert isinstance(data["rows"], list) and len(data["rows"]) == 4


def test_unevaluated_formula_rate_treated_as_missing(tmp_path):
    from openpyxl import load_workbook as _load
    base = tmp_path / "tender.xlsx"
    export_tender_to_excel(sample_compiled(), base)
    wb = _load(str(base))
    ws = wb["Tender Schedule"]
    for row in range(1, ws.max_row + 1):
        if str(ws.cell(row=row, column=1).value or "") == "1.1.1":
            # Programmatic save: no Excel calc cache, so data_only reads
            # None and the raw cell holds an untrusted formula string.
            ws.cell(row=row, column=5, value="=D10*2")
    formula = tmp_path / "formula.xlsx"
    wb.save(str(formula))
    summary = reconcile_tender_file(sample_compiled(), formula)
    row = [r for r in summary.rows if r.item_no == "1.1.1"][0]
    assert row.r_tender is None and row.variance is None
    assert "became-unpriced" in row.flags


def test_boolean_cells_never_parse_as_numbers(tmp_path):
    from openpyxl import load_workbook as _load
    base = tmp_path / "tender.xlsx"
    export_tender_to_excel(sample_compiled(), base)
    wb = _load(str(base))
    ws = wb["Tender Schedule"]
    for row in range(1, ws.max_row + 1):
        if str(ws.cell(row=row, column=1).value or "") == "2.1.1":
            ws.cell(row=row, column=4, value=True)
    bools = tmp_path / "bools.xlsx"
    wb.save(str(bools))
    summary = reconcile_tender_file(sample_compiled(), bools)
    row = [r for r in summary.rows if r.item_no == "2.1.1"][0]
    assert (row.q_tender, row.dq) == ("5", "0")
    assert "quantity-changed" not in row.flags


def test_reconciliation_rejects_bad_inputs(tmp_path):
    with pytest.raises(ValidationError):
        reconcile_tender_file("not-a-boq", tmp_path / "x.xlsx")
    with pytest.raises(ValidationError):
        reconcile_tender_file(sample_compiled(), tmp_path / "missing.xlsx")
    with pytest.raises(ValidationError):
        reconcile(sample_compiled(), "not-a-return")


# --- controllers ----------------------------------------------------------------------------

def test_tender_controllers_success_and_failure(tmp_path):
    from qs_engine.ui.viewmodels import ExportView
    ok_export = controllers.tender_export_controller(
        sample_compiled(), str(tmp_path / "tender.xlsx"),
        company_name="Test Corp", template_style="plain")
    assert isinstance(ok_export, ExportView)
    assert ok_export.path.endswith(".xlsx")
    bad_export = controllers.tender_export_controller(
        "not-a-boq", str(tmp_path / "x.xlsx"))
    assert type(bad_export).__name__ == "ErrorView"
    bad_style = controllers.tender_export_controller(
        sample_compiled(), str(tmp_path / "x.xlsx"),
        template_style="glossy")
    assert type(bad_style).__name__ == "ErrorView"
    path = controllers.tender_export_controller(
        sample_compiled(), str(tmp_path / "t.xlsx")).path
    summary = controllers.tender_reconciliation_controller(
        sample_compiled(), path)
    assert type(summary).__name__ == "ReconciliationSummary"
    assert summary.base_estimate == "2500"
    missing = controllers.tender_reconciliation_controller(
        sample_compiled(), str(tmp_path / "missing.xlsx"))
    assert type(missing).__name__ == "ErrorView"


# --- UI page wiring ------------------------------------------------------------------------------

def test_excel_page_dual_panel_wiring():
    import inspect as _inspect
    assert views.PAGES["excel"] is views.build_excel_page
    kinds = {key: kind for _g, items in views.NAV_GROUPS
             for key, _label, kind in items}
    assert kinds["excel"] == "page"
    source = _inspect.getsource(views.build_excel_page)
    assert "Export Settings & Action" in source
    assert "Reconciliation & Tender Review" in source
    assert "Run Variance Analysis" in source
    assert "RECON_COLUMNS" in source
    assert "Export BOQ" in source
    assert "session truth is never" in source
    assert "Unresolved" in source  # withheld lines labelled as such
    assert set(views.RECON_COLUMNS) == {
        "item", "description", "q_boq", "q_tender", "dq", "r_boq",
        "r_tender", "dr", "variance", "flags"}
