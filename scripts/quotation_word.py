"""Branded quotation Word generator (Phase 15).

Quotation DTO + company-profile metadata in, styled .docx out.
Accepts a DocumentPayload (or its plain-dict form) so this module stays
decoupled from qs_engine imports; amounts travel as exact strings.
"""

from datetime import date
from pathlib import Path

from docx import Document

from .document_base import (
    BaseBrandedDocument,
    apply_merge_tags,
    merge_mapping,
)

try:
    from qs_engine.core.errors import ValidationError
except Exception:  # pragma: no cover - standalone reuse outside the engine
    class ValidationError(Exception):
        """Fallback when reused outside the QS engine package."""


def _payload_dict(payload) -> dict:
    if hasattr(payload, "to_dict"):
        return payload.to_dict()
    if isinstance(payload, dict):
        return payload
    raise ValidationError(
        "Quotation needs a DocumentPayload (or its dict form); "
        f"got {type(payload).__name__}")


def _items(payload_dict: dict) -> list:
    items = []
    for section in payload_dict.get("sections", ()):
        for item in section.get("items", ()):
            items.append({
                "item_no": item.get("item_no", ""),
                "description": item.get("description", ""),
                "unit": item.get("unit", ""),
                "quantity": item.get("quantity", ""),
                "rate": item.get("rate"),
                "amount": item.get("amount"),
            })
    return items


def build_quotation_document(payload, profile: dict, template: str,
                             output_path, custom_template_path=None,
                             rate=None) -> str:
    """Compile a branded quotation .docx; returns the path as string."""
    data = _payload_dict(payload)
    base = BaseBrandedDocument(profile, template)
    title = data.get("document_title") or "Quotation"
    doc_no = (data.get("boq_ref") or data.get("compilation_id") or "")

    if template == "custom":
        if not custom_template_path or not Path(
                custom_template_path).is_file():
            raise ValidationError(
                "Custom Corporate Template needs an existing .docx "
                "template file; none was provided")
        doc = Document(str(custom_template_path))
        apply_merge_tags(doc, merge_mapping(data, profile or {}, title))
    else:
        doc = Document()
        base.build_header(doc, title)
        base.build_meta_table(doc, [
            ("Quotation No.:", doc_no or "—"),
            ("Date:", str(date.today().isoformat())),
            ("Project ref:", data.get("project_ref", "") or "—"),
            ("Template:", template),
        ])

    base.build_items_table(doc, _items(data))
    base.build_totals(doc, data.get("summary", {}) or {})
    base.build_rate_appendix(doc, rate)
    base.build_disclaimers(doc)
    base.build_footer(doc, f"Quotation No: {doc_no}" if doc_no else "")
    return BaseBrandedDocument.save(doc, output_path)
