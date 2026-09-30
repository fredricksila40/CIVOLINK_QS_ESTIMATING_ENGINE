"""Shared branded Word document base (Phase 15).

BaseBrandedDocument owns the header/identity block, styled BOQ tables,
totals, disclaimers, and footer shared by the quotation and invoice
generators. Presentation only: every amount arrives as a pre-formatted
string from the DTO layer; unresolved (None) amounts render as an
em-dash, never as zero.

Also hosts the merge-tag helper for custom corporate templates and a
minimal print-ready PDF summary writer (stdlib only — a faithful
Word-to-PDF conversion needs Word/LibreOffice, so the PDF is an honest
text summary, never claimed as a pixel-faithful conversion).
"""

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FALLBACK_ACCENT = "#004F9D"
CIVOLINK_SLATE = "#2F3A44"
MISSING = "\u2014"

#: Merge tags honoured inside user-supplied custom .docx templates.
CUSTOM_MERGE_TAGS = (
    "business_name",
    "document_title",
    "project_ref",
    "boq_ref",
    "grand_total",
    "document_date",
    "contact_phone",
    "email",
)


def accent_or_fallback(profile: dict) -> str:
    """Brand accent #RRGGBB from a profile, or CIVOLINK blue."""
    raw = str((profile or {}).get("brand_accent_color") or "").strip()
    if (len(raw) == 7 and raw.startswith("#")
            and all(ch in "0123456789ABCDEFabcdef" for ch in raw[1:])):
        return raw.upper()
    return FALLBACK_ACCENT


def _rgb(hex_color: str) -> RGBColor:
    hex_color = hex_color.lstrip("#")
    return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16),
                    int(hex_color[4:6], 16))


def _set_cell_fill(cell, hex_color: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), hex_color.replace("#", "").upper())
    cell._element.get_or_add_tcPr().append(shading)


def apply_merge_tags(doc: Document, mapping: dict) -> int:
    """Replace {{tag}} placeholders in body paragraphs and tables.

    Returns the number of replacements made. Unknown tags are left
    untouched so user templates never lose content silently.
    """
    count = 0
    targets = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                targets.extend(cell.paragraphs)
    for paragraph in targets:
        for tag in CUSTOM_MERGE_TAGS:
            token = "{{" + tag + "}}"
            if token in paragraph.text:
                value = str(mapping.get(tag, token))
                for run in paragraph.runs:
                    if token in run.text:
                        run.text = run.text.replace(token, value)
                        count += 1
                if token in paragraph.text:
                    paragraph.text = paragraph.text.replace(token, value)
                    count += 1
    return count


def _escape_pdf_text(text: str) -> str:
    return (text.replace("\\", "\\\\").replace("(", "\\(")
            .replace(")", "\\)"))


def write_summary_pdf(title: str, lines: list, output_path) -> str:
    """Write a minimal print-ready PDF text summary (stdlib only).

    A faithful Word-to-PDF conversion requires Word/LibreOffice, which
    cannot be assumed; this writer produces an honest A4 text summary
    carrying the same totals. Returns the output path as a string.
    """
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    body = [str(title), ""]
    body.extend(str(line) for line in lines)
    content = ["BT /F1 11 Tf 50 800 Td 14 TL"]
    for line in body:
        content.append(f"({_escape_pdf_text(line)}) Tj T*")
    content.append("ET")
    stream = "\n".join(content).encode("latin-1", errors="replace")
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        ("<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
         "/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"),
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(stream)} >>\nstream\n".encode(
            "latin-1") .decode("latin-1") + stream.decode("latin-1")
        + "\nendstream",
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = []
    for index, body_text in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf.extend(f"{index} 0 obj\n".encode("latin-1"))
        pdf.extend(body_text.encode("latin-1"))
        pdf.extend(b"\nendobj\n")
    xref = len(pdf)
    pdf.extend(f"xref\n0 {len(objects) + 1}\n".encode("latin-1"))
    pdf.extend(b"0000000000 65535 f \n")
    for offset in offsets:
        pdf.extend(f"{offset:010d} 00000 n \n".encode("latin-1"))
    pdf.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref}\n%%EOF".encode("latin-1"))
    output.write_bytes(bytes(pdf))
    return str(output)


class BaseBrandedDocument:
    """Shared branded-document builder over a company profile."""

    def __init__(self, profile: dict, template: str = "civolink"):
        self.profile = dict(profile or {})
        self.template = template
        self.accent = accent_or_fallback(self.profile)

    # -- building blocks -------------------------------------------------

    def header_fill(self) -> str:
        """Items-table header fill: brand accent, except the clean slate
        CIVOLINK tender style which stays neutral grey."""
        if self.template == "civolink":
            return CIVOLINK_SLATE
        return self.accent

    def build_header(self, doc: Document, document_title: str) -> None:
        profile = self.profile
        business = (profile.get("business_name") or "").strip()
        if self.template == "edmolt":
            brand_line = doc.add_paragraph()
            brand_line.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = brand_line.add_run(business or "Corporate Document")
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = _rgb(self.accent)
        else:
            brand_line = doc.add_paragraph()
            brand_line.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = brand_line.add_run(business or "Corporate Document")
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = _rgb(self.accent)
        contact_bits = [bit for bit in (
            (profile.get("physical_address") or "").strip(),
            " | ".join(bit for bit in (
                (profile.get("contact_phone") or "").strip(),
                (profile.get("email") or "").strip()) if bit),
        ) if bit]
        if contact_bits:
            contact = doc.add_paragraph()
            run = contact.add_run("\n".join(contact_bits))
            run.font.size = Pt(8)
        reg_tax = " | ".join(bit for bit in (
            (f"Reg/Tax ID: {profile.get('registration_tax_id').strip()}"
             if (profile.get("registration_tax_id") or "").strip() else ""),
        ) if bit)
        if reg_tax:
            line = doc.add_paragraph()
            run = line.add_run(reg_tax)
            run.font.size = Pt(8)
        logo = (profile.get("logo_path") or "").strip()
        if logo and Path(logo).is_file():
            try:
                doc.add_picture(logo, width=Cm(4.0))
            except Exception:
                pass  # a corrupt image must never kill an export
        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = title.add_run(document_title)
        run.font.size = Pt(22 if self.template == "edmolt" else 20)
        run.font.bold = True
        run.font.color.rgb = _rgb(self.accent)

    def build_meta_table(self, doc: Document, rows: list) -> None:
        table = doc.add_table(rows=len(rows), cols=2)
        table.style = "Table Grid"
        for index, (label, value) in enumerate(rows):
            label_cell, value_cell = table.rows[index].cells
            label_cell.text = str(label)
            value_cell.text = str(value)
            for run in label_cell.paragraphs[0].runs:
                run.font.bold = True
                run.font.size = Pt(9)
            for run in value_cell.paragraphs[0].runs:
                run.font.size = Pt(9)

    def build_items_table(self, doc: Document, items: list):
        """Styled BOQ items table; returns the table element."""
        table = doc.add_table(rows=1, cols=6)
        table.style = "Table Grid"
        headers = ("Item No.", "Description", "Unit",
                   "Quantity", "Rate", "Amount")
        fill = self.header_fill()
        for index, header in enumerate(headers):
            cell = table.rows[0].cells[index]
            cell.text = header
            _set_cell_fill(cell, fill)
            for run in cell.paragraphs[0].runs:
                run.font.bold = True
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(255, 255, 255)
        for item in items:
            row = table.add_row().cells
            row[0].text = str(item.get("item_no", ""))
            row[1].text = str(item.get("description", ""))
            row[2].text = str(item.get("unit", ""))
            row[3].text = str(item.get("quantity", ""))
            row[4].text = (str(item.get("rate"))
                           if item.get("rate") is not None else MISSING)
            row[5].text = (str(item.get("amount"))
                           if item.get("amount") is not None else MISSING)
            for cell in row:
                for run in cell.paragraphs[0].runs:
                    run.font.size = Pt(8)
        doc.add_paragraph()
        return table

    def build_totals(self, doc: Document, summary: dict) -> None:
        resolved = summary.get("resolved_subtotal", "")
        grand = summary.get("grand_total", None)
        table = doc.add_table(rows=2, cols=2)
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "Resolved subtotal"
        table.rows[0].cells[1].text = str(resolved)
        table.rows[1].cells[0].text = "Grand total"
        table.rows[1].cells[1].text = (
            str(grand) if grand is not None
            else "UNRESOLVED — see items marked " + MISSING)
        for row in table.rows:
            for run in row.cells[0].paragraphs[0].runs:
                run.font.bold = True
                run.font.size = Pt(9)
            for run in row.cells[1].paragraphs[0].runs:
                run.font.size = Pt(9)
        unresolved = summary.get("unresolved_items") or ()
        if unresolved:
            note = doc.add_paragraph()
            run = note.add_run(
                "Unresolved items: " + ", ".join(unresolved))
            run.font.size = Pt(8)

    def build_rate_appendix(self, doc: Document, rate) -> None:
        """Optional rate build-up appendix (RateSheetPayload or dict)."""
        if rate is None:
            return
        if hasattr(rate, "to_dict"):
            data = rate.to_dict()
        elif isinstance(rate, dict):
            data = rate
        else:
            return
        heading = doc.add_paragraph()
        run = heading.add_run("Rate Build-Up Appendix")
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = _rgb(self.accent)
        element = data.get("element", "")
        work_unit = data.get("work_unit", "")
        if element or work_unit:
            line = doc.add_paragraph()
            sub = line.add_run(f"{element} ({work_unit})".strip())
            sub.font.size = Pt(9)
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        for index, header in enumerate(
                ("Component", "Detail", "Cost / unit")):
            cell = table.rows[0].cells[index]
            cell.text = header
            _set_cell_fill(cell, self.header_fill())
            for run in cell.paragraphs[0].runs:
                run.font.bold = True
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(255, 255, 255)
        for component in data.get("components", ()):
            row = table.add_row().cells
            row[0].text = str(component.get("kind", ""))
            row[1].text = str(component.get("description", ""))
            row[2].text = str(component.get("cost_per_unit", ""))
            for cell in row:
                for run in cell.paragraphs[0].runs:
                    run.font.size = Pt(8)
        gross = doc.add_paragraph()
        run = gross.add_run(
            f"Gross rate: {data.get('gross_rate', '')}")
        run.font.size = Pt(9)
        run.font.bold = True
        doc.add_paragraph()

    def build_disclaimers(self, doc: Document) -> None:
        text = (self.profile.get("legal_disclaimers") or "").strip()
        if text:
            heading = doc.add_paragraph()
            run = heading.add_run("Legal / Disclaimers")
            run.font.size = Pt(10)
            run.font.bold = True
            run.font.color.rgb = _rgb(self.accent)
            body = doc.add_paragraph()
            run = body.add_run(text)
            run.font.size = Pt(8)
        governance = doc.add_paragraph()
        run = governance.add_run(
            "QS Engine = Source of Truth | Document = Presentation Surface")
        run.font.size = Pt(7)

    def build_footer(self, doc: Document, document_no: str = "") -> None:
        section = doc.sections[0]
        footer = (section.footer.paragraphs[0]
                  if section.footer.paragraphs
                  else section.footer.add_paragraph())
        parts = [self.profile.get("business_name", ""),
                 self.profile.get("contact_phone", ""),
                 self.profile.get("email", "")]
        if document_no:
            parts.append(document_no)
        footer.text = " | ".join(
            str(part).strip() for part in parts if str(part).strip())
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in footer.runs:
            run.font.size = Pt(7)

    @staticmethod
    def save(doc: Document, output_path) -> str:
        output = Path(output_path)
        output.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output))
        return str(output)


def merge_mapping(payload_dict: dict, profile: dict,
                  document_title: str) -> dict:
    """Merge-tag mapping for custom templates (all plain strings)."""
    summary = (payload_dict or {}).get("summary", {}) or {}
    grand = summary.get("grand_total")
    return {
        "business_name": (profile or {}).get("business_name", ""),
        "document_title": document_title,
        "project_ref": (payload_dict or {}).get("project_ref", ""),
        "boq_ref": (payload_dict or {}).get("boq_ref", ""),
        "grand_total": str(grand) if grand is not None else MISSING,
        "document_date": str(date.today().isoformat()),
        "contact_phone": (profile or {}).get("contact_phone", ""),
        "email": (profile or {}).get("email", ""),
    }
