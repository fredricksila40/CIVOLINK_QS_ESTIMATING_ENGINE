"""Tender reconciliation & variance engine (Phase 17).

Ingests a returned tender workbook (typically our tender schedule with
tenderer-edited Quantity / Rate cells) and computes exact row-by-row
variances against the canonical compiled BOQ — without touching
session truth. The session is never overwritten; reconciliation output
is a review DTO for estimator confirmation.

Reading discipline (openpyxl has no formula evaluator):
- Quantity / Rate cells are read data_only first (Excel-saved cached
  values), falling back to the raw cell when the cache is empty
  (programmatic edits). Unevaluated formula strings can never be
  trusted and count as missing — never guessed.
- Blank tender Quantity means "unchanged" (the tenderer did not touch
  it); blank tender Rate means Unpriced (None).
- Unpriced amounts stay None on both sides: variances are None
  whenever either side is unpriced — never zero-filled.
- Unresolved lines (the base total policy withholds their amounts
  from every total) stay open on the tender side as well: the
  exported Rate cell is the engine's own withheld line, never a
  tenderer price, so it never enters the tender total and never
  creates a variance.

All money math is exact Decimal; only the net variance % is rounded
(to 2dp, half-up) for display, and stays None on a zero base.
"""

from dataclasses import dataclass
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Optional, Tuple
import re

from openpyxl import load_workbook

from qs_engine.boq.models import CompiledBoq, STATUS_UNRESOLVED
from qs_engine.core.errors import ValidationError

TENDER_SHEET = "Tender Schedule"
TABLE_MARKER = "Item No"

#: Item-reference shape (dotted numerics like 1.1.1): structural rows
#: (section headers, totals, summary labels) never match this, so only
#: genuine unmatched item rows are counted, never our own chrome.
_ITEM_REF = re.compile(r"^[A-Za-z0-9._-]+$")

FLAG_QTY = "quantity-changed"
FLAG_RATE = "rate-changed"
FLAG_LUMP = "lump-changed"
FLAG_PRICED = "became-priced"
FLAG_UNPRICED = "became-unpriced"
FLAG_OPEN = "unpriced-open"
FLAG_UNRESOLVED = "unresolved-open"


@dataclass(frozen=True)
class TenderRow:
    """One ingested tender line keyed by compiled item number."""
    item_no: str
    quantity: Optional[Decimal]  # None = untouched (base applies)
    rate: Optional[Decimal]  # None = Unpriced
    amount: Optional[Decimal] = None  # tender lump (provisional rows)
    note: str = ""


@dataclass(frozen=True)
class TenderReturn:
    """Parsed tender file: matched rows plus honest skip counts."""
    rows: Tuple[TenderRow, ...]
    unmatched_rows: int = 0


@dataclass(frozen=True)
class RowVariance:
    """Exact per-item reconciliation (None = unpriced-open)."""
    item_no: str
    description: str
    unit: str
    q_boq: str
    r_boq: Optional[str]
    a_boq: Optional[str]
    q_tender: str
    r_tender: Optional[str]
    a_tender: Optional[str]
    dq: str
    dr: Optional[str]
    variance: Optional[str]
    modified: bool = False
    flags: Tuple[str, ...] = ()


@dataclass(frozen=True)
class ReconciliationSummary:
    """Review DTO: base vs tender roll-up plus change counts."""
    base_estimate: str
    tender_total: str
    tender_has_unpriced: bool = False
    variance_amount: str = "0"
    variance_pct: Optional[str] = None
    modified_count: int = 0
    quantity_changed_count: int = 0
    rate_changed_count: int = 0
    unpriced_count: int = 0
    unmatched_count: int = 0
    rows: Tuple[RowVariance, ...] = ()

    def to_dict(self) -> dict:
        """Plain-JSON dict (amounts already exact strings)."""
        import dataclasses

        def convert(value):
            if dataclasses.is_dataclass(value):
                return {f.name: convert(getattr(value, f.name))
                        for f in dataclasses.fields(value)}
            if isinstance(value, (tuple, list)):
                return [convert(v) for v in value]
            return value

        return convert(self)


def _numeric(value) -> Optional[Decimal]:
    """Cell value -> Decimal, or None for blanks/text/bools/formulas."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = str(value).strip()
    if not text or text.startswith("="):
        return None
    try:
        return Decimal(text)
    except Exception:
        return None  # "Unpriced", notes, and other text stay unpriced


def parse_tender_workbook(path, compiled: CompiledBoq) -> TenderReturn:
    """Ingest tender Quantity/Rate edits keyed by compiled item number."""
    if not isinstance(compiled, CompiledBoq):
        raise ValidationError(
            "Reconciliation needs a CompiledBoq; "
            f"got {type(compiled).__name__}")
    file_path = Path(path)
    if not file_path.is_file():
        raise ValidationError(
            f"Tender file not found: {file_path}")
    try:
        cached = load_workbook(str(file_path), data_only=True)
        raw = load_workbook(str(file_path), data_only=False)
    except Exception as exc:
        raise ValidationError(
            f"Could not read tender workbook {file_path.name}: "
            f"{exc}") from exc
    if TENDER_SHEET not in cached.sheetnames:
        raise ValidationError(
            f"Tender workbook has no {TENDER_SHEET!r} sheet — export a "
            f"tender schedule first, then reconcile the returned file")
    known = {item.item_no for section in compiled.sections
             for item in section.items}
    ws_cached = cached[TENDER_SHEET]
    ws_raw = raw[TENDER_SHEET]
    header_row = None
    for row in ws_cached.iter_rows(min_row=1, max_row=40, max_col=1,
                                   values_only=False):
        if str(row[0].value or "").strip() == TABLE_MARKER:
            header_row = row[0].row
            break
    if header_row is None:
        raise ValidationError(
            "No tender table found (missing 'Item No' header) — "
            "reconcile a file exported from this BOQ")
    rows: list[TenderRow] = []
    unmatched = 0
    for line in range(header_row + 1, ws_cached.max_row + 1):
        item_no = str(ws_cached.cell(row=line, column=1).value or "").strip()
        if not item_no:
            continue
        if item_no not in known:
            if _ITEM_REF.match(item_no) and any(
                    ch.isdigit() for ch in item_no):
                unmatched += 1
            continue
        quantity = _numeric(ws_cached.cell(row=line, column=4).value)
        if quantity is None:
            quantity = _numeric(ws_raw.cell(row=line, column=4).value)
        rate = _numeric(ws_cached.cell(row=line, column=5).value)
        if rate is None:
            rate = _numeric(ws_raw.cell(row=line, column=5).value)
        amount = _numeric(ws_cached.cell(row=line, column=6).value)
        if amount is None:
            amount = _numeric(ws_raw.cell(row=line, column=6).value)
        note = str(ws_cached.cell(row=line, column=7).value or "").strip()
        if note in ("", "Priced"):
            note = ""
        rows.append(TenderRow(item_no=item_no, quantity=quantity,
                              rate=rate, amount=amount, note=note))
    if not rows:
        raise ValidationError(
            "No tender rows matched compiled item numbers — reconcile "
            "the file exported from this BOQ")
    return TenderReturn(rows=tuple(rows), unmatched_rows=unmatched)


def reconcile(compiled: CompiledBoq,
              tender: TenderReturn) -> ReconciliationSummary:
    """Compute exact row-by-row variances (session truth untouched)."""
    if not isinstance(compiled, CompiledBoq):
        raise ValidationError(
            "Reconciliation needs a CompiledBoq; "
            f"got {type(compiled).__name__}")
    if not isinstance(tender, TenderReturn):
        raise ValidationError(
            "Reconciliation needs a TenderReturn; "
            f"got {type(tender).__name__}")
    by_item_no = {row.item_no: row for row in tender.rows}
    variances: list[RowVariance] = []
    tender_total = Decimal("0")
    tender_open = False
    for section in compiled.sections:
        for item in section.items:
            tender_row = by_item_no.get(item.item_no)
            q_tender = (tender_row.quantity
                        if tender_row is not None
                        and tender_row.quantity is not None
                        else item.quantity)
            dq = q_tender - item.quantity
            flags: list[str] = []
            if dq != 0:
                flags.append(FLAG_QTY)
            if item.pricing_kind == "provisional":
                # Lump sums carry no rate: the tender Amount cell is the
                # revised lump (blank means untouched); quantities never
                # scale a lump, so dq stays informational only.
                if (tender_row is not None
                        and tender_row.amount is not None):
                    a_tender = tender_row.amount
                else:
                    a_tender = item.amount
                r_tender = None
                dr = None
                variance = (a_tender - item.amount
                            if a_tender is not None
                            and item.amount is not None else None)
                if variance is not None and variance != 0:
                    flags.append(FLAG_LUMP)
            else:
                r_tender = (tender_row.rate
                            if tender_row is not None else item.rate)
                if item.status == STATUS_UNRESOLVED:
                    # Base total policy withholds this line from
                    # every total (unverified measurement rule). The
                    # tender mirror carries the rate but a blank
                    # amount — the engine's own withheld line,
                    # never a tenderer price: keep it out of the
                    # tender total exactly like unpriced lines.
                    a_tender = None
                else:
                    a_tender = (q_tender * r_tender
                                if r_tender is not None else None)
                if item.rate is None or r_tender is None:
                    dr = None
                else:
                    dr = r_tender - item.rate
                if item.amount is None or a_tender is None:
                    variance = None
                else:
                    variance = a_tender - item.amount
                if item.rate is None and r_tender is not None:
                    flags.append(FLAG_PRICED)
                elif item.rate is not None and r_tender is None:
                    flags.append(FLAG_UNPRICED)
                elif dr is not None and dr != 0:
                    flags.append(FLAG_RATE)
                if item.status == STATUS_UNRESOLVED:
                    flags.append(FLAG_UNRESOLVED)
            if a_tender is None:
                tender_open = True
                flags.append(FLAG_OPEN)
            else:
                tender_total += a_tender
            variances.append(RowVariance(
                item_no=item.item_no,
                description=item.description,
                unit=item.unit,
                q_boq=str(item.quantity),
                r_boq=(str(item.rate)
                       if item.rate is not None else None),
                a_boq=(str(item.amount)
                       if item.amount is not None else None),
                q_tender=str(q_tender),
                r_tender=(str(r_tender)
                          if r_tender is not None else None),
                a_tender=(str(a_tender)
                          if a_tender is not None else None),
                dq=str(dq),
                dr=(str(dr) if dr is not None else None),
                variance=(str(variance)
                          if variance is not None else None),
                modified=any(
                    flag not in (FLAG_OPEN, FLAG_UNRESOLVED)
                    for flag in flags),
                flags=tuple(flags),
            ))
    base = compiled.resolved_subtotal
    difference = tender_total - base
    pct = (None if base == 0 else str(
        (difference / base * 100).quantize(
            Decimal("0.01"), rounding=ROUND_HALF_UP)))
    return ReconciliationSummary(
        base_estimate=str(base),
        tender_total=str(tender_total),
        tender_has_unpriced=tender_open,
        variance_amount=str(difference),
        variance_pct=pct,
        modified_count=sum(1 for row in variances if row.modified),
        quantity_changed_count=sum(
            1 for row in variances if FLAG_QTY in row.flags),
        rate_changed_count=sum(
            1 for row in variances
            if FLAG_RATE in row.flags or FLAG_PRICED in row.flags
            or FLAG_UNPRICED in row.flags or FLAG_LUMP in row.flags),
        unpriced_count=sum(
            1 for row in variances if FLAG_OPEN in row.flags),
        unmatched_count=tender.unmatched_rows,
        rows=tuple(variances),
    )


def reconcile_tender_file(compiled: CompiledBoq,
                          path) -> ReconciliationSummary:
    """Parse a returned tender file and reconcile it in one step."""
    return reconcile(compiled, parse_tender_workbook(path, compiled))
