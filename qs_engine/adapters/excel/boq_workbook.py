"""BOQ -> Excel workbook exporter (Phase 11).

Read-only presentation sink: the compiled Phase 10 BOQ is the truth;
this module only renders it. Reconciliation design (per consolidated
brief): Amount (F) carries a native Excel formula for priced measured
items, Compiled Amount (G) carries the deterministic compiled Decimal,
Check (H) carries =F-G. Provisional lumps render as values (no formula
exists); unpriced renders as UNPRICED text (never zero).

PRECISION NOTE: .xlsx numerics are IEEE-754 doubles — openpyxl coerces
Decimal values to float on write. Display stays exact for money via
#,##0.00 formatting; the compiled Decimal truth lives in qs_engine,
not in this file. Styling is neutral grey: EDMOLT brand assets/colours
are frozen in another product and must never leak in here.
"""

from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from qs_engine import __version__ as QS_ENGINE_VERSION
from qs_engine.boq.models import (
    STATUS_PROVISIONAL,
    STATUS_UNPRICED,
    CompiledBoq,
)
from qs_engine.core.errors import ValidationError

GOVERNANCE_BANNER = (
    "QS Engine = Source of Truth | Excel = Export / Review Surface"
)

HEADERS = ("Item No", "Description", "Quantity", "Unit", "Rate",
           "Amount", "Compiled Amount", "Check")

QTY_FORMAT = "#,##0.000"
MONEY_FORMAT = "#,##0.00"

HEADER_FILL = PatternFill("solid", start_color="D9D9D9")
SUBSECTION_FILL = PatternFill("solid", start_color="EAEAEA")
TOTAL_FILL = PatternFill("solid", start_color="BFBFBF")
BANNER_FILL = PatternFill("solid", start_color="F2F2F2")

HEADER_FONT = Font(bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14)
BOLD_FONT = Font(bold=True)
ITALIC_GREY = Font(italic=True, size=9, color="595959")

THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"))

COLUMN_WIDTHS = {"A": 12, "B": 48, "C": 14, "D": 10,
                 "E": 16, "F": 18, "G": 18, "H": 14}

UNPRICED_TEXT = "UNPRICED"
UNRESOLVED_TEXT = "UNRESOLVED"
PROVISIONAL_SUFFIX = " (Provisional Sum)"


def sanitize_sheet_name(name: str, used: set) -> str:
    """Excel-safe unique sheet title (<=31 chars, no []:*?/\\)."""
    cleaned = "".join("_" if ch in '[]:*?/\\' else ch
                      for ch in str(name)).strip() or "Sheet"
    cleaned = cleaned[:31].rstrip()
    candidate, suffix = cleaned, 2
    while candidate in used:
        tail = f"-{suffix}"
        candidate = (cleaned[:31 - len(tail)].rstrip() + tail)
        suffix += 1
    used.add(candidate)
    return candidate


def _style_header_row(ws, row: int, fill=HEADER_FILL) -> None:
    for col in range(1, 9):
        cell = ws.cell(row=row, column=col)
        cell.font = HEADER_FONT
        cell.fill = fill
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="center",
                                   vertical="center", wrap_text=True)


def _money(cell, value: Decimal) -> None:
    cell.value = value
    cell.number_format = MONEY_FORMAT
    cell.border = THIN_BORDER


def _qty(cell, value: Decimal) -> None:
    cell.value = value
    cell.number_format = QTY_FORMAT
    cell.border = THIN_BORDER


def _prepare_sheet(ws, title: str) -> None:
    for letter, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A3"
    ws.print_title_rows = "2:2"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:H1")
    for index, header in enumerate(HEADERS, start=1):
        ws.cell(row=2, column=index, value=header)
    _style_header_row(ws, 2)
    ws.row_dimensions[2].height = 30


def _write_item_row(ws, row: int, item) -> None:
    """Render one compiled item per the pricing-kind contract."""
    ws.cell(row=row, column=1, value=item.item_no).border = THIN_BORDER
    description = item.description
    if item.status == STATUS_PROVISIONAL:
        description += PROVISIONAL_SUFFIX
    desc_cell = ws.cell(row=row, column=2, value=description)
    desc_cell.border = THIN_BORDER
    desc_cell.alignment = Alignment(vertical="center", wrap_text=True)
    _qty(ws.cell(row=row, column=3), item.quantity)
    unit_cell = ws.cell(row=row, column=4, value=item.unit)
    unit_cell.border = THIN_BORDER
    unit_cell.alignment = Alignment(horizontal="center")

    amount_cell = ws.cell(row=row, column=6)
    compiled_cell = ws.cell(row=row, column=7)
    check_cell = ws.cell(row=row, column=8)
    if item.status == STATUS_UNPRICED:
        amount_cell.value = UNPRICED_TEXT
        amount_cell.border = THIN_BORDER
        compiled_cell.value = UNPRICED_TEXT
        compiled_cell.border = THIN_BORDER
        check_cell.border = THIN_BORDER
        rate_cell = ws.cell(row=row, column=5)
        rate_cell.border = THIN_BORDER
        return
    if item.pricing_kind == "provisional":
        ws.cell(row=row, column=5).border = THIN_BORDER  # no rate: lump sum
        _money(amount_cell, item.amount)
        _money(compiled_cell, item.amount)
        check_cell.value = f"=F{row}-G{row}"
        check_cell.number_format = MONEY_FORMAT
        check_cell.border = THIN_BORDER
        return
    # Priced measured item: formula, compiled value, reconciliation.
    _money(ws.cell(row=row, column=5), item.rate)
    amount_cell.value = f"=C{row}*E{row}"
    amount_cell.number_format = MONEY_FORMAT
    amount_cell.border = THIN_BORDER
    _money(compiled_cell, item.amount)
    check_cell.value = f"=F{row}-G{row}"
    check_cell.number_format = MONEY_FORMAT
    check_cell.border = THIN_BORDER


def _write_subsection_header(ws, row: int, code: str, title: str) -> None:
    """Label-only grouping row (never referenced by any total)."""
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
    cell = ws.cell(row=row, column=1, value=f"{code} — {title}")
    cell.font = BOLD_FONT
    cell.fill = SUBSECTION_FILL
    for col in range(1, 9):
        ws.cell(row=row, column=col).border = THIN_BORDER
        ws.cell(row=row, column=col).fill = SUBSECTION_FILL


def _write_total_row(ws, row: int, label: str,
                     total) -> None:
    """Section total: compiled value or UNRESOLVED text, never a partial
    formula sum (a formula would silently contradict compiled None)."""
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    label_cell = ws.cell(row=row, column=1, value=label)
    label_cell.font = BOLD_FONT
    label_cell.fill = TOTAL_FILL
    label_cell.alignment = Alignment(horizontal="right")
    for col in range(1, 9):
        ws.cell(row=row, column=col).fill = TOTAL_FILL
        ws.cell(row=row, column=col).border = THIN_BORDER
        ws.cell(row=row, column=col).font = BOLD_FONT
    value_cell = ws.cell(row=row, column=7)
    if total is None:
        value_cell.value = UNRESOLVED_TEXT
    else:
        _money(value_cell, total)
        value_cell.fill = TOTAL_FILL
        value_cell.font = BOLD_FONT


def _write_section_sheet(wb, section, used: set):
    ws = wb.create_sheet(sanitize_sheet_name(section.code, used))
    _prepare_sheet(ws, f"SECTION {section.code} — {section.title}")
    from qs_engine.boq.models import CompiledSubsection
    by_number = {item.item_no: item for item in section.items}
    row = 3
    subs = section.subsections or (
        # Defensive: hand-built sections predating subsection detail
        # group all items under a single header row.
        (CompiledSubsection(code=section.code, title=section.title,
                            item_nos=tuple(by_number)),),
    )
    for sub in subs:
        _write_subsection_header(ws, row, sub.code, sub.title)
        row += 1
        for number in sub.item_nos:
            _write_item_row(ws, row, by_number[number])
            row += 1
    _write_total_row(ws, row, "SECTION TOTAL", section.total)
    return ws


def _cover_row(ws, row: int, label: str, value) -> None:
    label_cell = ws.cell(row=row, column=1, value=label)
    label_cell.font = BOLD_FONT
    value_cell = ws.cell(row=row, column=2, value=value)
    if isinstance(value, Decimal):
        value_cell.number_format = MONEY_FORMAT
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)


def _write_cover_sheet(wb, compiled: CompiledBoq, project_name: str,
                       exported_on: str, engine_version: str):
    ws = wb.active
    ws.title = "Cover"
    for letter, width in (("A", 24), ("B", 30), ("C", 30), ("D", 30)):
        ws.column_dimensions[letter].width = width
    ws["A1"] = "BILL OF QUANTITIES"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:D1")
    ws["A2"] = GOVERNANCE_BANNER
    ws["A2"].font = ITALIC_GREY
    ws["A2"].fill = BANNER_FILL
    ws.merge_cells("A2:D2")

    total_items = sum(len(s.items) for s in compiled.sections)
    status = ("COMPLETE" if not compiled.has_unresolved
              else f"{len(compiled.unresolved_items)} UNRESOLVED ITEM(S)")
    rules = sorted({i.measurement_ref for s in compiled.sections
                    for i in s.items if i.measurement_ref})
    rows = [
        ("Project reference", compiled.project_ref),
        ("Project name", project_name or "—"),
        ("BOQ reference", compiled.boq_ref),
        ("Compilation status", status),
        ("Grand total", (compiled.total if compiled.total is not None
                         else UNRESOLVED_TEXT)),
        ("Sections", len(compiled.sections)),
        ("Items", total_items),
        ("Unresolved items",
         ", ".join(compiled.unresolved_items) or "None"),
        ("Measurement rules cited", "; ".join(rules) or "None cited"),
        ("Engine", f"CIVOLINK QS Engine {engine_version}"),
    ]
    if exported_on:
        rows.append(("Exported on", exported_on))
    row = 4
    for label, value in rows:
        _cover_row(ws, row, label, value)
        if label == "Grand total" and isinstance(value, Decimal):
            ws.cell(row=row, column=2).font = BOLD_FONT
        row += 1
    ws.freeze_panes = "A4"
    return ws


def _write_summary_sheet(wb, compiled: CompiledBoq):
    ws = wb.create_sheet("Summary")
    for letter, width in (("A", 14), ("B", 48), ("C", 16), ("D", 20)):
        ws.column_dimensions[letter].width = width
    ws["A1"] = "GRAND SUMMARY"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:D1")
    for index, header in enumerate(
            ("Section", "Title", "Status", "Amount"), start=1):
        cell = ws.cell(row=2, column=index, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="center")
    # Compiled values (never partial formula sums: SUM() would skip
    # UNPRICED text and silently contradict a compiled None total).
    row = 3
    for section in compiled.sections:
        ws.cell(row=row, column=1, value=section.code).border = THIN_BORDER
        ws.cell(row=row, column=2,
                value=section.title).border = THIN_BORDER
        status_cell = ws.cell(
            row=row, column=3,
            value="COMPLETE" if not section.has_unresolved else UNRESOLVED_TEXT)
        status_cell.border = THIN_BORDER
        status_cell.alignment = Alignment(horizontal="center")
        amount_cell = ws.cell(row=row, column=4)
        if section.total is None:
            amount_cell.value = UNRESOLVED_TEXT
        else:
            _money(amount_cell, section.total)
        row += 1
    grand_label = ws.cell(row=row, column=1, value="GRAND TOTAL")
    grand_label.font = BOLD_FONT
    grand_label.fill = TOTAL_FILL
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    for col in range(1, 5):
        ws.cell(row=row, column=col).fill = TOTAL_FILL
        ws.cell(row=row, column=col).border = THIN_BORDER
        ws.cell(row=row, column=col).font = BOLD_FONT
    grand_cell = ws.cell(row=row, column=4)
    if compiled.total is None:
        grand_cell.value = UNRESOLVED_TEXT
    else:
        _money(grand_cell, compiled.total)
        grand_cell.fill = TOTAL_FILL
        grand_cell.font = BOLD_FONT
    ws.freeze_panes = "A3"
    ws.print_title_rows = "2:2"
    return ws


def export_boq_to_excel(compiled: CompiledBoq, output_path,
                        *, project_name: str = "",
                        exported_on: str = "",
                        engine_version: str = None) -> str:
    """Export a compiled BOQ to a professional .xlsx workbook.

    Cover (audit metadata + governance banner), Summary (section roll-up),
    then one bill sheet per section with F/G/H reconciliation columns.
    Returns the output path as a string.
    """
    if not isinstance(compiled, CompiledBoq):
        raise ValidationError(
            f"export_boq_to_excel needs a CompiledBoq; "
            f"got {type(compiled).__name__}")
    if engine_version is None:
        engine_version = QS_ENGINE_VERSION
    wb = Workbook()
    _write_cover_sheet(wb, compiled, project_name, exported_on,
                       engine_version)
    _write_summary_sheet(wb, compiled)
    used = {"Cover", "Summary"}
    for section in compiled.sections:
        _write_section_sheet(wb, section, used)
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output))
    return str(output)
