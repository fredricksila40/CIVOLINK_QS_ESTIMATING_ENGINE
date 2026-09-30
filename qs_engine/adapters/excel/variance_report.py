"""Site Material Variance & Store Reconciliation Report writer (Phase 6.2).

openpyxl is a presentation dependency and lives ONLY inside
``qs_engine/adapters/`` (enforced by the import-boundary suite). This
adapter renders a ``StoreReconciliationSummary`` into a formatted
``.xlsx`` workbook: a summary sheet with the head-line metrics and a
per-material variance sheet whose Status column carries colour-coded
alert fills (green / amber / red).
"""

from pathlib import Path

from qs_engine.core.errors import ValidationError
from qs_engine.reconciliation.models import (
    STATUS_CRITICAL_OVERRUN,
    STATUS_NORMAL,
    STATUS_WARNING,
    StoreReconciliationSummary,
)

#: Status -> (fill, font) alert colours.
STATUS_STYLES = {
    STATUS_NORMAL: ("C6EFCE", "006100"),
    STATUS_WARNING: ("FFEB9C", "9C6500"),
    STATUS_CRITICAL_OVERRUN: ("FFC7CE", "9C0006"),
}

REPORT_TITLE = "Site Material Variance & Store Reconciliation Report"

SUMMARY_HEADERS = ("Metric", "Value")
SUMMARY_ROWS = (
    ("Total Materials Tracked", "total_materials_tracked"),
    ("Theoretical Budget Value", "theoretical_budget_value"),
    ("Actual Issued Value", "actual_issued_value"),
    ("Overall Material Variance", "overall_variance_value"),
    ("Warning Threshold (%)", "warning_threshold_pct"),
    ("Critical Multiplier", "critical_multiplier"),
)

VARIANCE_HEADERS = (
    "Material Code", "Material Name", "Unit", "Theoretical Qty",
    "Actual Issued", "Store Balance", "Variance Qty", "Variance %",
    "Unit Rate", "Variance Value", "Status", "Wastage %",
)


def write_variance_report(path, summary: StoreReconciliationSummary,
                          title: str = REPORT_TITLE) -> Path:
    """Write the reconciliation summary to a formatted .xlsx workbook.

    Returns the written path. Raises ValidationError when the summary
    has no rows (there is nothing to report) or the path is not .xlsx.
    """
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        raise ValidationError(
            f"Variance report must be written to .xlsx; got {path.suffix!r}")
    if not summary.rows:
        raise ValidationError(
            "No reconciliation rows to export — import a store log and "
            "compile a BOQ first")

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()

    # -- summary sheet -----------------------------------------------------
    sheet = workbook.active
    sheet.title = "Summary"
    sheet["A1"] = title
    sheet["A1"].font = Font(bold=True, size=14, color="004F9D")
    sheet.merge_cells("A1:B1")
    sheet["A2"] = ("Theoretical QS allowances vs actual site consumption; "
                   "exact Decimal arithmetic, no float rounding.")
    sheet["A2"].font = Font(italic=True, color="555D6B")
    sheet.merge_cells("A2:B2")

    header_fill = PatternFill("solid", fgColor="004F9D")
    header_font = Font(bold=True, color="FFFFFF")
    for column, header in enumerate(SUMMARY_HEADERS, start=1):
        cell = sheet.cell(row=4, column=column, value=header)
        cell.fill = header_fill
        cell.font = header_font
    for offset, (label, attr) in enumerate(SUMMARY_ROWS, start=5):
        sheet.cell(row=offset, column=1, value=label)
        value = getattr(summary, attr)
        sheet.cell(row=offset, column=2,
                   value="—" if value is None else str(value))
    sheet.column_dimensions["A"].width = 30
    sheet.column_dimensions["B"].width = 24

    # -- variance sheet ----------------------------------------------------
    grid = workbook.create_sheet("Material Variance")
    for column, header in enumerate(VARIANCE_HEADERS, start=1):
        cell = grid.cell(row=1, column=column, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    for index, row in enumerate(summary.rows, start=2):
        values = (
            row.code, row.name, row.unit, str(row.theoretical_qty),
            str(row.actual_issued), str(row.store_balance),
            str(row.variance_qty),
            ("" if row.variance_pct is None else str(row.variance_pct)),
            ("" if row.unit_rate is None else str(row.unit_rate)),
            ("" if row.variance_value is None else str(row.variance_value)),
            row.status,
            ("" if row.wastage_pct is None else str(row.wastage_pct)),
        )
        for column, value in enumerate(values, start=1):
            grid.cell(row=index, column=column, value=value)
        fill_hex, font_hex = STATUS_STYLES[row.status]
        status_cell = grid.cell(row=index, column=11)
        status_cell.fill = PatternFill("solid", fgColor=fill_hex)
        status_cell.font = Font(bold=True, color=font_hex)
    widths = (16, 24, 8, 14, 13, 13, 12, 11, 10, 14, 18, 10)
    for column, width in enumerate(widths, start=1):
        grid.column_dimensions[get_column_letter(column)].width = width
    grid.freeze_panes = "A2"

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path
